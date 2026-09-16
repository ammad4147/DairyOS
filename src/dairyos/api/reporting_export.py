"""Binary exporters for the governed DairyOS Reporting dataset."""

from __future__ import annotations

import csv
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape, portrait
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


PDF_MARGIN = 10 * mm
PDF_LANDSCAPE_COLUMN_THRESHOLD = 6


OPERATOR_HEADINGS: dict[str, str] = {
    "animal_id": "Animal ID",
    "ear_tag": "Ear Tag",
    "rfid": "RFID",
    "date_of_birth": "Date of Birth",
    "date_of_acquisition": "Date of Acquisition",
    "dam_id": "Dam ID",
    "sire_id": "Sire ID",
    "lifecycle_status": "Lifecycle Status",
    "is_currently_milking": "Currently Milking",
    "milking_frequency": "Milking Frequency",
    "production_date": "Date",
    "operational_date": "Date",
    "event_date": "Event Date",
    "recorded_at": "Recorded At",
    "sample_date": "Sample Date",
    "herd_total_label": "Herd Group",
    "total_yield": "Total Milk (L)",
    "morning_yield": "Morning (L)",
    "afternoon_yield": "Afternoon (L)",
    "evening_yield": "Evening (L)",
    "selected_session": "Milking Session",
    "selected_session_yield": "Session Milk (L)",
    "quantity_liters": "Quantity (L)",
    "amount": "Amount, PKR",
    "feed_cost": "Feed Cost (PKR)",
    "total_herd_feed_cost_per_day": "Daily Herd Feed Cost (PKR)",
    "feed_cost_per_litre_today": "Feed Cost / Litre (PKR)",
    "cost_per_head_day": "Cost / Head / Day (PKR)",
    "price_per_kg": "Price / kg (PKR)",
    "record_id": "Record ID",
    "report_id": "Report",
    "record_count": "Records",
    "batch_id": "Batch ID", "feed_source_lot": "Feed Lot Reference", "breed_code": "Breed Code", "production_phase_dim": "Production Phase (DIM)",
    "somatic_cell_count": "SCC (cells/mL)", "antibiotic_residue_status": "Antibiotic Status", "cooling_chain_break": "Cooling Chain Break", "adulteration_test_result": "Adulteration Test",
    "iso_17025_ref": "ISO 17025 Reference", "analyst_id": "Analyst ID", "retention_until": "Retention Until",
}


def _heading(value: str) -> str:
    key = str(value or "").strip()
    if key in OPERATOR_HEADINGS:
        return OPERATOR_HEADINGS[key]
    text = re.sub(r"[_\-]+", " ", key).strip()
    return " ".join(word.upper() if word.upper() in {"ID", "AI", "PD", "TMR", "COP", "COML", "PKR", "RFID"} else word.capitalize() for word in text.split())


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return " · ".join(f"{_heading(str(key))}: {_cell(item)}" for key, item in value.items())
    if isinstance(value, (list, tuple)):
        return ", ".join(_cell(item) for item in value)
    return str(value)


def _xlsx_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    return _cell(value)


def _pdf_page_size(columns: list[str]) -> tuple[float, float]:
    """Return the governed A4 orientation for a Reporting table."""
    return landscape(A4) if len(columns) > PDF_LANDSCAPE_COLUMN_THRESHOLD else portrait(A4)


def csv_bytes(columns: list[str], rows: list[dict[str, Any]]) -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\r\n")
    writer.writerow([_heading(column) for column in columns])
    for row in rows:
        writer.writerow([_cell(row.get(column)) for column in columns])
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def xlsx_bytes(title: str, columns: list[str], rows: list[dict[str, Any]], summary: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append([_heading(column) for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        sheet.append([_xlsx_value(row.get(column)) for column in columns])
    sheet.freeze_panes = "A2"
    if columns:
        sheet.auto_filter.ref = sheet.dimensions
    for index, cells in enumerate(sheet.columns, start=1):
        width = min(42, max(11, max(len(_cell(cell.value)) for cell in cells) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    meta = workbook.create_sheet("Summary")
    meta.append(["DairyOS Report", title])
    meta.append(["Record Count", len(rows)])
    for key, value in summary.items():
        meta.append([_heading(str(key)), _cell(value)])
    meta.column_dimensions["A"].width = 28
    meta.column_dimensions["B"].width = 72
    for cell in meta[1]:
        cell.font = Font(bold=True)
    for row in meta.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_bytes(title: str, columns: list[str], rows: list[dict[str, Any]], summary: dict[str, Any]) -> bytes:
    page_size = _pdf_page_size(columns)
    wide = page_size[0] > page_size[1]
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=PDF_MARGIN,
        rightMargin=PDF_MARGIN,
        topMargin=PDF_MARGIN,
        bottomMargin=PDF_MARGIN,
        title=title,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(f"DairyOS — {title}", styles["Title"]), Spacer(1, 3*mm)]
    if summary:
        summary_data = [[Paragraph(_heading(str(key)), styles["BodyText"]), Paragraph(_cell(value), styles["BodyText"])] for key, value in summary.items()]
        summary_table = Table(summary_data, colWidths=[42*mm, None])
        summary_table.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
        story.extend([summary_table, Spacer(1, 3*mm)])

    if columns:
        header_style = styles["BodyText"].clone("ReportHeader"); header_style.fontName = "Helvetica-Bold"; header_style.fontSize = 8.5; header_style.leading = 10
        body_style = styles["BodyText"].clone("ReportBody"); body_style.fontSize = 8.5 if not wide else 8; body_style.leading = 10
        data = [[Paragraph(_heading(column), header_style) for column in columns]]
        for row in rows:
            data.append([Paragraph(_cell(row.get(column)), body_style) for column in columns])
        available_width = page_size[0] - (2 * PDF_MARGIN)
        weights = [max(7, min(24, len(_heading(column)))) for column in columns]
        total_weight = sum(weights) or 1
        widths = [available_width * weight / total_weight for weight in weights]
        table = Table(data, repeatRows=1, colWidths=widths)
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("GRID",(0,0),(-1,-1),0.25,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3)]))
        story.append(table)
    else:
        story.append(Paragraph("No records match the selected report controls.", styles["BodyText"]))

    document.build(story)
    return output.getvalue()
